#!/usr/bin/env python3
"""Fill column B with a "Recall" sentence built from the founder note in column A.

Usage: python3 fill_recall.py input.xlsx [output.xlsx]

Only column B of the active sheet is written; every other cell is left as-is.
Prints a one-line summary: rows filled and rows flagged for review.
"""
import datetime
import re
import sys
from pathlib import Path

import openpyxl

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

# First M/YY or M/YYYY in the cell, not glued to other digits.
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})/(\d{4}|\d{2})(?!\d)")


def sentence_for(value):
    """Return (text, flagged) for one column-A value."""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return "REVIEW: column A is stored as an Excel date, not text (year may be wrong)", True
    if not isinstance(value, str):
        return f"REVIEW: column A is not text ({value!r})", True

    text = " ".join(value.lower().split())
    # "not raising" first: "not raising now" also contains "raising now".
    if "not raising" in text:
        clause = "you were not raising"
    elif "raising now" in text:
        clause = "you were considering a round"
    elif "early" in text:
        clause = "it was early for us"
    else:
        clause = None

    m = DATE_RE.search(value)
    if not m and clause is None:
        return "REVIEW: no date and no \"early\"/\"not raising\"/\"raising now\"", True
    if not m:
        return "REVIEW: no date", True
    if clause is None:
        return "REVIEW: no \"early\", \"not raising\" or \"raising now\"", True

    month = int(m.group(1))
    year = int(m.group(2))
    if len(m.group(2)) == 2:
        year += 2000
    if not 1 <= month <= 12:
        return f"REVIEW: invalid month in \"{m.group(0)}\"", True

    # Months from the last 11 months show the bare month name; older ones get the
    # year suffix (e.g. on 10/1/2026, 11/25 -> "November", 10/25 -> "October ’25").
    today = datetime.date.today()
    months_ago = (today.year * 12 + today.month) - (year * 12 + month)
    if year <= 2024:
        when = str(year)
    elif months_ago >= 12:
        when = f"{MONTHS[month - 1]} ’{year % 100:02d}"
    else:
        when = MONTHS[month - 1]

    # "raising now" puts a comma after the date and has no "We last connected ... but" form.
    if clause == "you were considering a round":
        return f"When we connected in {when}, {clause}.  Recall ", False
    if year <= 2024:
        return f"We last connected in {year} but {clause}.  Recall ", False
    return f"When we connected in {when} {clause}.  Recall ", False


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(f"{src.stem}_filled{src.suffix}")

    wb = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    ws = wb.active

    filled, flagged = [], []
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=1).value
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        text, is_flagged = sentence_for(value)
        ws.cell(row=row, column=2).value = text
        (flagged if is_flagged else filled).append(row)

    wb.save(dst)
    print(f"Saved {dst}")
    print(f"{len(filled)} rows filled; flagged rows: "
          f"{', '.join(map(str, flagged)) if flagged else 'none'}")


if __name__ == "__main__":
    main()
